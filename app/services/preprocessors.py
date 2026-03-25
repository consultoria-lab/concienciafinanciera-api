import base64
import io

import openpyxl
import pymupdf
from langsmith import traceable
from openai import OpenAI

from app.config import get_settings


def pdf_pages_to_base64_images(file_bytes: bytes) -> list[str]:
    """Convierte cada pagina del PDF a una imagen base64 para vision del LLM."""
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    images = []
    for page in doc:
        # Renderizar pagina como imagen PNG a 200 DPI (buen balance calidad/tamano)
        pix = page.get_pixmap(dpi=200)
        img_bytes = pix.tobytes("png")
        b64 = base64.b64encode(img_bytes).decode("utf-8")
        images.append(f"data:image/png;base64,{b64}")
    doc.close()
    return images


def encode_image_to_base64(file_bytes: bytes, content_type: str) -> str:
    """Codifica imagen a base64 data URI para GPT-4o-mini vision."""
    b64 = base64.b64encode(file_bytes).decode("utf-8")
    return f"data:{content_type};base64,{b64}"


def extract_text_from_excel(file_bytes: bytes) -> str:
    """Lee un Excel y lo convierte a texto tabular para el LLM."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):  # Ignorar filas completamente vacias
                rows.append(" | ".join(cells))
        if rows:
            if len(wb.sheetnames) > 1:
                parts.append(f"--- Hoja: {sheet_name} ---")
            parts.append("\n".join(rows))
    wb.close()
    return "\n\n".join(parts).strip()


@traceable(run_type="llm", name="Whisper Transcription")
def transcribe_audio(file_bytes: bytes, filename: str) -> str:
    """Transcribe audio usando Whisper API de OpenAI."""
    settings = get_settings()
    client = OpenAI(api_key=settings.openai_api_key)
    audio_file = io.BytesIO(file_bytes)
    audio_file.name = filename
    transcript = client.audio.transcriptions.create(
        model=settings.whisper_model,
        file=audio_file,
        language="es",
    )
    return transcript.text
